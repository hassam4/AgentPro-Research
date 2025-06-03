import os
from pptx import Presentation
from pptx.util import Inches


def create_ppt(summaries, output_path="summary.pptx") -> str:
    prs = Presentation()
    for item in summaries:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = item["title"]
        body = slide.shapes.placeholders[1]
        tf = body.text_frame
        tf.text = item["summary"]
    prs.save(output_path)
    return os.path.abspath(output_path)
